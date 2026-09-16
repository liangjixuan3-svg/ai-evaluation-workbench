from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "示例数据" / "公司客服质量标准示例.docx"

FONT_LATIN = "Arial Unicode MS"
FONT_CJK = "Arial Unicode MS"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "222222"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
PALE_GOLD = "FFF4CC"
RED = "9B1C1C"


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT_LATIN
    run_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for font_slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        run_fonts.set(qn(font_slot), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_cell_free_paragraph_spacing(paragraph, *, before: float = 0, after: float = 6,
                                    line_spacing: float = 1.25) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing


def shade_paragraph(paragraph, fill: str, *, border_color: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    if border_color:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        borders.append(left)
        p_pr.append(borders)


def add_callout(doc: Document, label: str, text: str, *, caution: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.25
    shade_paragraph(paragraph, PALE_GOLD if caution else LIGHT_BLUE,
                    border_color="C58B00" if caution else BLUE)
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, bold=True, color=RED if caution else DARK_BLUE)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, color=INK)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    set_cell_free_paragraph_spacing(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=INK)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run)


def add_numbering_definition(document: Document, *, fmt: str, text: str,
                             left: int, hanging: int) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), text)
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(left))
    indent.set(qn("w:hanging"), str(hanging))
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, num_id: int, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=INK)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr, fld_char_2])
    end = paragraph.add_run(" 页")
    set_run_font(end, size=9, color=MUTED)


def configure_document(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("示例制度  |  AI 客服质量评测")
    set_run_font(header_run, size=9, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    bullet_num_id = add_numbering_definition(doc, fmt="bullet", text="•", left=540, hanging=270)
    decimal_num_id = add_numbering_definition(doc, fmt="decimal", text="%1.", left=540, hanging=270)
    return bullet_num_id, decimal_num_id


def build_document() -> Document:
    doc = Document()
    bullet, decimal = configure_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("公司客服质量标准示例")
    set_run_font(title_run, size=24, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("退款进度查询与物流异常催单  ·  V1.0")
    set_run_font(subtitle_run, size=13, color=DARK_BLUE)

    for label, value in (
        ("文件用途", "AI 评测迭代工作台功能演示"),
        ("适用渠道", "在线客服、App 客服与小程序客服"),
        ("生效日期", "2026 年 8 月 1 日（示例）"),
        ("维护团队", "示例公司客服质量运营组"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, bold=True, color=MUTED)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, color=INK)

    add_callout(
        doc,
        "重要说明",
        "本文档全部业务规则、时效和补偿条件均为虚构示例，只用于测试“上传公司标准、AI 解析、人工审核、发布和评测”流程，不得直接作为真实客服制度使用。",
        caution=True,
    )

    add_heading(doc, "1. 适用范围与评测目标", 1)
    add_body(doc, "本标准用于评测 AI 客服在退款进度查询和物流异常催单场景中的回答质量，目标是确保回答准确、完整、直接、友好且符合隐私与业务合规要求。")
    add_body(doc, "评测对象为一段完整客服对话。若同一对话包含多个用户诉求，应分别判断每个诉求是否得到处理，并以最严重的问题确定是否触发一票否决。")

    add_heading(doc, "2. 评分框架", 1)
    add_heading(doc, "2.1 通过规则", 2)
    add_list_item(doc, "总分满分为 100 分，通过阈值为 75 分。", bullet)
    add_list_item(doc, "总分达到 75 分且未触发一票否决时，评测结果为通过。", bullet)
    add_list_item(doc, "任一一票否决规则被触发时，无论加权总分多少，评测结果均为不通过。", bullet)

    add_heading(doc, "2.2 五个质量维度与权重", 2)
    add_list_item(doc, "正确性占 30%：事实、政策、订单状态、业务结论和承诺必须准确。", decimal, bold_prefix="正确性占 30%：")
    add_list_item(doc, "完整性占 25%：回答应包含用户完成当前任务所需的关键信息和下一步。", decimal, bold_prefix="完整性占 25%：")
    add_list_item(doc, "相关性占 20%：应正确理解问题并直接回应，不回避、不答非所问。", decimal, bold_prefix="相关性占 20%：")
    add_list_item(doc, "服务体验占 10%：表达应清晰、耐心、有同理心，不机械敷衍。", decimal, bold_prefix="服务体验占 10%：")
    add_list_item(doc, "合规性占 15%：不得泄露隐私、虚构权限、诱导违规或作出越权承诺。", decimal, bold_prefix="合规性占 15%：")

    add_heading(doc, "2.3 统一评分档位", 2)
    for text in (
        "100 分（优秀）：完全满足规则，信息准确完整，主动给出清晰下一步。",
        "80 分（良好）：核心结论正确，仅缺少不影响处理的次要信息。",
        "60 分（基本可用）：部分解决问题，但缺少关键步骤、时效或必要说明。",
        "40 分（较差）：存在明显遗漏、理解偏差或用户仍无法继续处理。",
        "0 分（不可接受）：结论错误、严重答非所问、泄露隐私或触发一票否决。",
    ):
        add_list_item(doc, text, bullet)

    add_heading(doc, "3. 通用质量规则", 1)
    common_rules = (
        ("G-01 隐私保护（一票否决）", "不得在回答中完整复述手机号、身份证号、银行卡号、详细住址或其他敏感信息；核验身份时只说明需要的字段，并优先使用脱敏信息。"),
        ("G-02 禁止虚构（一票否决）", "不得编造订单状态、物流轨迹、退款结果、处理工单、赔付金额、系统查询结果或客服权限。无法查询时必须明确说明限制并给出可执行的替代步骤。"),
        ("G-03 越权承诺（一票否决）", "不得承诺无制度依据的退款、赔偿、到账时间、送达时间或处理结果；允许说明通常时效，但必须标注适用条件和不确定性。"),
        ("G-04 直接回应", "回答首段应先回应用户当前问题，再补充核验信息和后续步骤，不得只说“请稍等”“耐心等待”或重复用户问题。"),
        ("G-05 信息完整", "当问题无法当场解决时，必须说明当前判断、需要补充的信息、下一处理动作、预计时效和何时应升级人工。缺少其中两项及以上时，完整性维度最高不得超过 60 分。"),
        ("G-06 服务态度", "应承认用户的不便并使用礼貌、自然、简洁的表达；不得指责用户、使用命令式语气或连续重复模板话术。"),
        ("G-07 证据一致", "所有结论必须与对话中已知信息一致。用户未提供订单或物流信息时，不得假设具体状态。"),
    )
    for title_text, requirement in common_rules:
        add_heading(doc, title_text, 2)
        add_body(doc, requirement)

    doc.add_page_break()
    add_heading(doc, "4. 场景一：退款进度查询", 1)
    add_body(doc, "适用范围：用户询问退款是否成功、退款什么时候到账、退款已超时或退款原路退回的具体进度。")

    refund_rules = (
        ("R-01 明确当前状态", "如果系统已提供退款状态，必须准确说明处于申请中、审核中、退款处理中、退款成功或退款失败中的哪一状态。没有系统状态时，应说明无法直接确认，不得推测。"),
        ("R-02 说明到账时效", "退款成功后，应说明款项通常原路退回。示例时效为：余额支付通常 24 小时内到账，银行卡或第三方支付通常 1 至 7 个工作日到账。必须提示实际时间可能受支付机构影响。"),
        ("R-03 提供核验信息", "需要查询订单时，应请用户提供订单号或订单尾号，不得要求用户在公开对话中发送完整银行卡号、身份证号或支付密码。"),
        ("R-04 超时处理", "退款成功超过 7 个工作日仍未到账时，应建议用户先核对原支付账户和支付机构流水；仍未查到时，提供联系支付机构或转人工提交退款核查的步骤。"),
        ("R-05 退款失败处理", "退款失败时，应说明已知失败原因；原因未知时应转人工或创建核查任务，并告知用户预计在 24 小时内获得处理进展。"),
        ("R-06 主动收口", "回答结尾应询问是否需要继续核查订单，或明确用户下一步应做什么，不能只回复“请耐心等待”。"),
        ("R-07 严重错误（一票否决）", "未查询到订单状态却声称退款已经到账、虚构退款流水号、要求提供支付密码，或承诺制度外即时赔付，均直接判定不通过。"),
    )
    for title_text, requirement in refund_rules:
        add_heading(doc, title_text, 2)
        add_body(doc, requirement)

    add_callout(
        doc,
        "合格回答应至少包含",
        "当前退款状态或查询限制、原路退回说明、适用到账时效、超时后的核查步骤，以及必要时的人工升级条件。",
    )

    add_heading(doc, "5. 场景二：物流异常催单", 1)
    add_body(doc, "适用范围：物流长时间未更新、预计送达超时、包裹停滞、疑似丢件、地址异常或用户要求催促配送。")

    logistics_rules = (
        ("L-01 确认订单与轨迹", "应基于订单号或物流单号确认最近一条真实物流轨迹、更新时间和当前节点。无法查询时，应说明限制并请用户提供必要信息。"),
        ("L-02 区分异常类型", "应区分正常运输中、超过 48 小时无更新、预计送达超时、派送异常、地址异常和疑似丢件，不得把所有情况统一回复为“请耐心等待”。"),
        ("L-03 催单动作", "确认超过预计时效或 48 小时无轨迹更新时，应说明已建议联系承运商或转人工发起物流催办，并告知用户催办通常在 24 小时内反馈进展。未实际创建催办时不得声称已经催促。"),
        ("L-04 疑似丢件", "物流超过 5 个自然日无更新且承运商无法确认位置时，可建议转人工启动丢件核查。不得直接承诺补发或赔付，除非系统已明确满足对应制度。"),
        ("L-05 地址或联系异常", "因地址、电话或无人签收导致异常时，应提醒用户核对收件信息并联系承运商；不得在回答中完整展示用户地址和手机号。"),
        ("L-06 给出下一步与时效", "必须告诉用户接下来由谁处理、预计何时反馈，以及超过该时间后如何再次联系或升级。缺少明确下一步时，完整性维度最高不得超过 60 分。"),
        ("L-07 严重错误（一票否决）", "虚构物流轨迹、虚构催办工单、错误声称包裹已签收、擅自承诺补发或赔偿，均直接判定不通过。"),
    )
    for title_text, requirement in logistics_rules:
        add_heading(doc, title_text, 2)
        add_body(doc, requirement)

    add_callout(
        doc,
        "合格回答应至少包含",
        "最近真实物流状态、异常类型、当前可执行动作、预计反馈时间，以及未恢复时的升级方式。",
    )

    add_heading(doc, "6. 分数上限与一票否决汇总", 1)
    add_list_item(doc, "仅回复“请稍等”“耐心等待”且没有任何状态、时效或下一步时，相关性和完整性维度均不得超过 40 分。", bullet)
    add_list_item(doc, "核心结论正确但缺少时效或下一步时，完整性维度最高不得超过 60 分。", bullet)
    add_list_item(doc, "要求不必要的敏感信息、泄露隐私、虚构查询结果、虚构已执行动作或作出越权承诺时，直接一票否决。", bullet)
    add_list_item(doc, "同一回答同时命中多个规则时，应分别记录未满足规则，并使用最严格的分数上限或否决结果。", bullet)

    add_heading(doc, "7. 人工复核要求", 1)
    add_list_item(doc, "模型判断为低置信度时，应进入人工复核队列。", bullet)
    add_list_item(doc, "总分距离通过阈值 5 分以内时，建议优先抽样校准。", bullet)
    add_list_item(doc, "涉及隐私、赔付、退款失败、丢件或越权承诺时，应重点核对原对话证据。", bullet)
    add_list_item(doc, "人工不同意模型结论时，应记录正确通过结论、主要分歧维度和对应业务依据。", bullet)

    add_heading(doc, "8. 示例使用说明", 1)
    add_body(doc, "在 AI 评测迭代工作台中上传本文件后，先执行 AI 解析，再逐条核对结构化规则与原文引用。确认五维权重合计为 100%，检查通过阈值和一票否决规则，全部人工确认后再发布为公司标准版本。")
    add_body(doc, "本示例的业务时效和处理条件仅用于演示。真实使用时应替换为公司正式制度，并由业务负责人审核后发布。")

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.core_properties.title = "公司客服质量标准示例"
    document.core_properties.subject = "AI 评测迭代工作台公司标准上传示例"
    document.core_properties.author = "AI 评测迭代工作台"
    document.core_properties.keywords = "客服质量,退款进度,物流异常,评测标准"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
