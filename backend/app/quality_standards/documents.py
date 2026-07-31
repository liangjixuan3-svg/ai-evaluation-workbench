from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_DOCUMENT_SIZE = 20 * 1024 * 1024


@dataclass(frozen=True)
class SourceSection:
    locator: str
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    sections: list[SourceSection]
    sha256: str


def extract_document(filename: str, content: bytes) -> ExtractedDocument:
    """提取 DOCX 段落或 PDF 页面中的可复制文字。"""
    if len(content) > MAX_DOCUMENT_SIZE:
        raise ValueError("单个文件不能超过 20 MB，请拆分后再上传")

    suffix = Path(filename).suffix.lower()
    if suffix not in {".docx", ".pdf"}:
        raise ValueError("仅支持 .docx 和 .pdf 文件，请检查上传文件")

    if suffix == ".docx":
        _ensure_docx_content(content)
        sections = _extract_docx_sections(content)
    else:
        _ensure_pdf_content(content)
        sections = _extract_pdf_sections(content)

    if not sections:
        raise ValueError("未提取到可复制的文本，请上传包含文本的文档；扫描版 PDF 暂不支持")

    return ExtractedDocument(
        text="\n".join(section.text for section in sections),
        sections=sections,
        sha256=sha256(content).hexdigest(),
    )


def _ensure_docx_content(content: bytes) -> None:
    if not content.startswith(b"PK"):
        raise ValueError("文件扩展名与实际内容不匹配，请上传正确的 .docx 文件")


def _ensure_pdf_content(content: bytes) -> None:
    if not content.startswith(b"%PDF-"):
        raise ValueError("文件扩展名与实际内容不匹配，请上传正确的 .pdf 文件")


def _extract_docx_sections(content: bytes) -> list[SourceSection]:
    try:
        document = Document(BytesIO(content))
    except (PackageNotFoundError, ValueError) as error:
        raise ValueError("DOCX 文档无法读取，请确认文件未损坏后重新上传") from error

    return [
        SourceSection(locator=f"第 {index} 段", text=text)
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if (text := paragraph.text.strip())
    ]


def _extract_pdf_sections(content: bytes) -> list[SourceSection]:
    try:
        reader = PdfReader(BytesIO(content))
        page_texts = [page.extract_text().strip() for page in reader.pages]
    except PdfReadError as error:
        raise ValueError("PDF 文档无法读取，请确认文件未损坏后重新上传") from error

    return [
        SourceSection(locator=f"第 {index} 页", text=text)
        for index, text in enumerate(page_texts, start=1)
        if text
    ]
